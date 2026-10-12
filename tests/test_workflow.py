"""The launch path, end to end over HTTP: score -> submit -> review -> approve -> publish -> parent."""

from io import BytesIO

from django.core.cache import cache
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from accounts.models import User
from ops.models import SmsMessage
from records.models import Student
from results.models import ClassCycle, ResultSheet, Score, SubjectSheet
from tests.factories import PASSWORD, SchoolWorld


def xlsx(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    buffer.name = "upload.xlsx"
    return buffer


class PublishJourneyTests(TestCase):
    def setUp(self):
        cache.clear()
        self.w = SchoolWorld("journey", students=3)

    def _enter_scores(self, user, assignment):
        self.client.force_login(user)
        for offset, enrollment in enumerate(self.w.enrollments):
            for component, value in ((self.w.ca, 30 - offset), (self.w.exam, 50 - offset)):
                response = self.client.post(
                    self.w.url(f"app/scores/{assignment.id}/save/"),
                    {"enrollment": enrollment.id, "component": component.id, "value": str(value)},
                )
                self.assertEqual(response.json()["ok"], True)
        self.client.post(self.w.url(f"app/scores/{assignment.id}/submit/"))

    def test_full_journey(self):
        self._enter_scores(self.w.teacher, self.w.assignment)
        self._enter_scores(self.w.other_teacher, self.w.maths_assignment)
        self.assertEqual(SubjectSheet.objects.filter(state=SubjectSheet.State.SUBMITTED).count(), 2)

        # Submitted sheets are locked for the teacher.
        locked = self.client.post(
            self.w.url(f"app/scores/{self.w.maths_assignment.id}/save/"),
            {"enrollment": self.w.enrollments[0].id, "component": self.w.ca.id, "value": "1"},
        )
        self.assertEqual(locked.status_code, 400)

        self.client.force_login(self.w.class_teacher)
        review_url = self.w.url(f"app/review/{self.w.classroom.id}/")
        self.assertEqual(self.client.get(review_url).status_code, 200)
        self.client.post(review_url, {"action": "review"})
        cycle = ClassCycle.objects.get(school_class=self.w.classroom, term=self.w.term)
        self.assertEqual(cycle.state, ClassCycle.State.REVIEWED)

        # A class teacher cannot approve their own review.
        self.client.post(review_url, {"action": "approve"})
        cycle.refresh_from_db()
        self.assertEqual(cycle.state, ClassCycle.State.REVIEWED)

        self.client.force_login(self.w.head)
        self.client.post(review_url, {"action": "approve", f"hs_{self.w.enrollments[0].id}": "Well done."})
        cycle.refresh_from_db()
        self.assertEqual(cycle.state, ClassCycle.State.APPROVED)

        self.client.force_login(self.w.admin)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(self.w.url("app/publish/"), {"scope": "school", "sms": "on"})
        sheets = ResultSheet.objects.filter(term=self.w.term)
        self.assertTrue(all(s.state == ResultSheet.State.PUBLISHED for s in sheets))
        self.assertTrue(all(s.result_code and s.snapshot for s in sheets))
        self.assertEqual(SmsMessage.objects.filter(status=SmsMessage.Status.SENT).count(), 1)

        sheet = sheets.get(enrollment=self.w.enrollments[0])
        self.client.force_login(self.w.parent)
        page = self.client.get(self.w.url(f"portal/result/{sheet.id}/"))
        self.assertContains(page, "Well done.")
        pdf = self.client.get(self.w.url(f"portal/result/{sheet.id}/pdf/"))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(b"".join(pdf.streaming_content).startswith(b"%PDF"))

        self.client.logout()
        verify = self.client.get(self.w.url(f"verify/{sheet.result_code}/"))
        self.assertContains(verify, self.w.student.surname)

    def test_publish_is_all_or_nothing(self):
        self.w.score_everything()
        ClassCycle.objects.create(school_class=self.w.classroom, term=self.w.term, state=ClassCycle.State.APPROVED)
        from schools.models import SchoolClass

        SchoolClass.objects.create(school=self.w.school, level=self.w.level, arm="B", name="JSS 1 B")
        self.client.force_login(self.w.admin)
        self.client.post(self.w.url("app/publish/"), {"scope": "school"})
        # JSS 1 B was never approved, so JSS 1 A must not go out alone.
        self.assertFalse(ResultSheet.objects.filter(state=ResultSheet.State.PUBLISHED).exists())

    def test_unicode_names_publish_to_pdf(self):
        Student.objects.filter(pk=self.w.student.pk).update(surname="Ọlábísí", first_name="Ẹ̀mẹ́kà")
        self.w.publish()
        from results.services.pdf import generate_result_pdf

        sheet = ResultSheet.objects.get(enrollment=self.w.enrollments[0])
        self.assertTrue(generate_result_pdf(sheet).name.endswith(".pdf"))


class ExcelRoundTripTests(TestCase):
    def setUp(self):
        cache.clear()
        self.w = SchoolWorld("excel", students=2)
        self.client.force_login(self.w.teacher)
        self.url = self.w.url(f"app/scores/{self.w.assignment.id}/excel/")

    def test_download_then_upload(self):
        download = self.client.get(self.url)
        self.assertEqual(download.status_code, 200)
        book = load_workbook(BytesIO(download.content))
        rows = list(book.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][:2], ("admission_number", "name"))

        upload = xlsx(
            [
                ["admission_number", "name", "CA", "Exam"],
                [self.w.enrollments[0].student.admission_number, "x", 35, 55],
                [self.w.enrollments[1].student.admission_number.lower(), "y", "AB", "nan"],
                ["NOPE/999", "z", 10, 10],
            ]
        )
        response = self.client.post(self.url, {"file": upload}, follow=True)
        self.assertContains(response, "Saved 3 scores")
        self.assertEqual(Score.objects.filter(marker=Score.Marker.ABSENT).count(), 1)
        self.assertEqual(Score.objects.get(enrollment=self.w.enrollments[0], component=self.w.exam).value, 55)

    def test_garbage_file_is_a_message(self):
        junk = BytesIO(b"not a spreadsheet")
        junk.name = "junk.xlsx"
        response = self.client.post(self.url, {"file": junk}, follow=True)
        self.assertContains(response, "not a valid Excel workbook")


class StudentImportTests(TestCase):
    def setUp(self):
        cache.clear()
        self.w = SchoolWorld("import", students=0)
        self.client.force_login(self.w.admin)

    def test_import_creates_students_parents_and_enrollments(self):
        upload = xlsx(
            [
                ["admission_number", "surname", "first_name", "sex", "class", "arm", "parent_phone", "status"],
                ["IMP/1", "Okoro", "Amaka", "F", "JSS 1", "A", "08051234567", "active"],
                ["IMP/2", "Bello", "Musa", "m", "JSS 1", "A", "", "bogus"],
                ["IMP/3", "Eze", "Ada", "X", "JSS 1", "A", "", ""],
            ]
        )
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(self.w.url("app/students/import/"), {"file": upload})
        self.assertEqual(self.w.school.students.count(), 2)
        self.assertEqual(Student.objects.get(admission_number="IMP/2").status, Student.Status.ACTIVE)
        parent = User.objects.get(phone="08051234567")
        self.assertEqual(parent.role, User.Role.PARENT)
        self.assertFalse(parent.check_password("parent1234"))
        self.assertTrue(parent.must_change_password)
        invite = SmsMessage.objects.get(phone="08051234567")
        self.assertIn("********", invite.body)

    def test_import_never_links_a_child_to_a_staff_phone(self):
        upload = xlsx(
            [
                ["admission_number", "surname", "first_name", "sex", "class", "arm", "parent_phone"],
                ["IMP/9", "Okoro", "Amaka", "F", "JSS 1", "A", self.w.head.phone],
            ]
        )
        self.client.post(self.w.url("app/students/import/"), {"file": upload})
        self.assertIsNone(Student.objects.get(admission_number="IMP/9").parent_user)


class PasswordChangeTests(TestCase):
    def setUp(self):
        cache.clear()
        self.w = SchoolWorld("pw", students=0)
        self.w.teacher.must_change_password = True
        self.w.teacher.save()

    def test_temporary_password_must_be_changed_first(self):
        self.client.post("/login/", {"identifier": self.w.teacher.phone, "password": PASSWORD})
        blocked = self.client.get(self.w.url("app/scores/"))
        self.assertEqual(blocked.status_code, 302)
        self.assertIn("/password/", blocked["Location"])

        weak = self.client.post("/password/", {"current": PASSWORD, "new": "12345678", "confirm": "12345678"})
        self.assertEqual(weak.status_code, 200)
        self.client.post(
            "/password/", {"current": PASSWORD, "new": "Tidy-Blue-Lamp-42", "confirm": "Tidy-Blue-Lamp-42"}
        )
        self.w.teacher.refresh_from_db()
        self.assertFalse(self.w.teacher.must_change_password)
        self.assertEqual(self.client.get(self.w.url("app/scores/")).status_code, 200)


class HealthTests(TestCase):
    def test_healthz(self):
        response = self.client.get("/healthz/")
        self.assertEqual(response.json(), {"status": "ok"})


class ReviewFormTests(TestCase):
    def test_saved_rating_is_shown_selected(self):
        from schools.models import BehaviouralTrait

        cache.clear()
        w = SchoolWorld("rate", students=1)
        trait = BehaviouralTrait.objects.create(school=w.school, name="Punctuality", domain="affective")
        w.score_everything()
        self.client.force_login(w.class_teacher)
        url = w.url(f"app/review/{w.classroom.id}/")
        self.client.get(url)
        enrollment = w.enrollments[0]
        self.client.post(url, {"action": "save_review", f"trait_{enrollment.id}_{trait.id}": "4"})
        self.assertContains(self.client.get(url), '<option value="4" selected>')

    def test_branded_404(self):
        response = self.client.get("/definitely-not-a-school/")
        self.assertEqual(response.status_code, 404)
        self.assertTemplateUsed(response, "404.html")
