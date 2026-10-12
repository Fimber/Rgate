import secrets
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook

from accounts.models import User
from common.parsing import open_workbook
from records.models import Enrollment, Student
from schools.models import SchoolClass

TEMPLATE_HEADERS = [
    "admission_number",
    "surname",
    "first_name",
    "other_names",
    "sex",
    "date_of_birth",
    "class",
    "arm",
    "parent_name",
    "parent_phone",
    "parent_email",
    "admission_date",
    "status",
]


def student_template():
    book = Workbook()
    sheet = book.active
    sheet.title = "Students"
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(
        [
            "GFC/001",
            "Okoro",
            "Amaka",
            "Chioma",
            "F",
            "2014-03-12",
            "JSS 1",
            "A",
            "Ngozi Okoro",
            "08055555555",
            "parent@example.com",
            "2024-09-08",
            "active",
        ]
    )
    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


def import_students(school, uploaded, session, term):
    book = open_workbook(uploaded)
    sheet = book.active
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [str(value or "").strip() for value in first_row]
    created = updated = 0
    errors = []
    with transaction.atomic():
        for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            data = dict(zip(headers, row))
            number = str(data.get("admission_number") or "").strip()
            if not number:
                errors.append({"row": index, "error": "Admission number is required"})
                continue
            sex = str(data.get("sex") or "").strip().upper()[:1]
            if sex not in {"M", "F"}:
                errors.append({"row": index, "error": "Sex must be M or F"})
                continue
            student, was_created = Student.objects.update_or_create(
                school=school,
                admission_number=number,
                defaults={
                    "surname": str(data.get("surname") or "").strip(),
                    "first_name": str(data.get("first_name") or "").strip(),
                    "other_names": str(data.get("other_names") or "").strip(),
                    "sex": sex,
                    "parent_name": str(data.get("parent_name") or "").strip(),
                    "parent_phone": User.objects.normalize_phone(data.get("parent_phone") or "")
                    if data.get("parent_phone")
                    else "",
                    "parent_email": str(data.get("parent_email") or "").strip(),
                    "status": _status(data.get("status")),
                },
            )
            created += int(was_created)
            updated += int(not was_created)
            level_name = str(data.get("class") or "").strip()
            arm = str(data.get("arm") or "A").strip()
            if level_name and term:
                school_class = SchoolClass.objects.filter(
                    school=school, level__name__iexact=level_name, arm__iexact=arm
                ).first()
                if not school_class:
                    errors.append({"row": index, "error": f"Unknown class {level_name} {arm}"})
                    continue
                Enrollment.objects.get_or_create(
                    student=student,
                    term=term,
                    defaults={"school_class": school_class, "session": session, "status": Enrollment.Status.ACTIVE},
                )
                _ensure_parent(school, student)
    return {"created": created, "updated": updated, "errors": errors}


def _status(raw):
    value = str(raw or "").strip().lower()
    return value if value in Student.Status.values else Student.Status.ACTIVE


def _ensure_parent(school, student):
    if not student.parent_phone or student.parent_user_id:
        return
    parent = User.objects.filter(phone=student.parent_phone).first()
    if parent is not None and (parent.role != User.Role.PARENT or parent.school_id != school.id):
        return  # never attach a child to a staff account or to another school's user
    if parent is None:
        from results.services.notify import send_parent_invite

        password = secrets.token_hex(4)
        parent = User.objects.create_user(
            phone=student.parent_phone,
            password=password,
            first_name=student.parent_name.split(" ")[0] if student.parent_name else "Parent",
            last_name=student.surname,
            role=User.Role.PARENT,
            school=school,
            email=student.parent_email,
            must_change_password=True,
        )
        transaction.on_commit(lambda: send_parent_invite(school, parent, password))
    student.parent_user = parent
    student.save(update_fields=["parent_user"])


def enroll_active_students(school, session, term):
    """Carry each active student into the selected term using their latest class."""
    count = 0
    for student in school.students.filter(status=Student.Status.ACTIVE):
        latest = student.enrollments.select_related("school_class").order_by("-id").first()
        if latest is None:
            continue
        _, created = Enrollment.objects.get_or_create(
            student=student,
            term=term,
            defaults={"school_class": latest.school_class, "session": session, "status": Enrollment.Status.ACTIVE},
        )
        count += int(created)
    return count
